"""Safe Phase 3 asset, import, onboarding, and profile services."""

from __future__ import annotations

import csv
from dataclasses import dataclass
import hashlib
from io import BytesIO, StringIO
import ipaddress
import json
from pathlib import Path
import re
import socket
from typing import Any, Callable
from urllib.parse import urlparse
import zipfile

from defusedxml import ElementTree as SafeET
from docx import Document
from openpyxl import load_workbook
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, sessionmaker

from .accounts import AccountService, AuthorizationError, Principal
from .database import set_workspace_context
from .models import (
    BrandProfileVersion,
    ImportJob,
    OnboardingInterview,
    OnboardingTurn,
    Product,
    WorkspaceAsset,
    new_id,
    utc_now,
)
from .storage import LocalStorage, StorageError


class Phase3Error(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class ParsedUpload:
    content: bytes
    asset_type: str
    mime_type: str
    metadata: dict[str, Any]
    extracted_text: str | None = None
    product_rows: tuple[dict[str, str], ...] = ()


MAX_ZIP_FILES = 500


def _extension(filename: str) -> str:
    return Path(filename).suffix.casefold()


def _zip_guard(content: bytes, max_archive_bytes: int) -> zipfile.ZipFile:
    if len(content) > max_archive_bytes:
        raise Phase3Error("archive_too_large", "Compressed archive exceeds the safety limit.")
    try:
        archive = zipfile.ZipFile(BytesIO(content))
        members = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise Phase3Error("malformed_archive", "The uploaded archive is malformed.") from exc
    if len(members) > MAX_ZIP_FILES:
        archive.close()
        raise Phase3Error("archive_too_many_files", "The archive contains too many files.")
    total = 0
    for member in members:
        name = Path(member.filename)
        if name.is_absolute() or ".." in name.parts:
            archive.close()
            raise Phase3Error("archive_path_traversal", "The archive contains an unsafe path.")
        total += member.file_size
        if member.compress_size and member.file_size / member.compress_size > 1000:
            archive.close()
            raise Phase3Error("decompression_bomb", "The archive compression ratio is unsafe.")
    if total > max_archive_bytes:
        archive.close()
        raise Phase3Error("archive_expands_too_large", "The archive expands beyond the safety limit.")
    return archive


def _svg(content: bytes) -> bytes:
    if len(content) > 5_000_000:
        raise Phase3Error("file_too_large", "SVG exceeds the safety limit.")
    lowered = content.decode("utf-8", errors="strict").casefold()
    if any(marker in lowered for marker in ("<!doctype", "<!entity", "<script", "foreignobject")):
        raise Phase3Error("unsafe_svg", "SVG scripts, entities, and foreign content are not allowed.")
    try:
        root = SafeET.fromstring(content)
    except Exception as exc:
        raise Phase3Error("malformed_svg", "The SVG document is malformed.") from exc
    for element in root.iter():
        if any(str(key).casefold().startswith("on") for key in element.attrib):
            raise Phase3Error("unsafe_svg", "SVG event handlers are not allowed.")
        for key, value in element.attrib.items():
            if str(key).casefold().endswith("href") and str(value).casefold().startswith(
                ("http:", "https:", "file:", "data:")
            ):
                raise Phase3Error("unsafe_svg", "SVG external references are not allowed.")
    return SafeET.tostring(root, encoding="utf-8")


def _image(filename: str, content: bytes) -> tuple[dict[str, Any], bytes]:
    try:
        with Image.open(BytesIO(content)) as image:
            image.verify()
        with Image.open(BytesIO(content)) as image:
            width, height = image.size
            orientation = image.getexif().get(274)
            mime_type = Image.MIME.get(image.format)
    except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
        raise Phase3Error("invalid_image", "The image is malformed or unsafe.") from exc
    if mime_type not in {"image/png", "image/jpeg", "image/webp"}:
        raise Phase3Error("unsupported_image", "Only PNG, JPEG, and WebP images are supported.")
    return {
        "width": width,
        "height": height,
        "orientation": str(orientation) if orientation else "normal",
        "format": image.format.casefold(),
    }, content


def _csv_rows(content: bytes, max_rows: int) -> tuple[list[str], tuple[dict[str, str], ...]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise Phase3Error("invalid_encoding", "CSV files must use UTF-8 encoding.") from exc
    if len(text) > 25_000_000:
        raise Phase3Error("file_too_large", "CSV content exceeds the safety limit.")
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        raise Phase3Error("missing_headers", "CSV requires a header row.")
    rows: list[dict[str, str]] = []
    for index, row in enumerate(reader):
        if index >= max_rows:
            raise Phase3Error("too_many_rows", "The import contains too many rows.")
        rows.append({str(key).strip(): (value or "").strip() for key, value in row.items() if key})
    return [str(value).strip() for value in reader.fieldnames], tuple(rows)


def _normalize_header(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().casefold())


def _product_rows(headers: list[str], rows: tuple[dict[str, str], ...]) -> tuple[dict[str, str], ...]:
    aliases = {
        "sku": {"sku", "code", "product code", "كود", "رمز المنتج", "الرمز"},
        "name": {"name", "product name", "اسم المنتج", "الاسم"},
        "description": {"description", "الوصف"},
        "price": {"price", "السعر"},
        "availability": {"availability", "available", "التوفر", "متاح"},
    }
    mapping: dict[str, str] = {}
    for header in headers:
        normalized = _normalize_header(header)
        for target, values in aliases.items():
            if normalized in values:
                mapping[target] = header
    if "name" not in mapping and "sku" not in mapping:
        raise Phase3Error("unmapped_columns", "The import needs a product name or SKU column.")
    normalized_rows: list[dict[str, str]] = []
    for row in rows:
        name = row.get(mapping.get("name", ""), "").strip()
        sku = row.get(mapping.get("sku", ""), "").strip() or name.casefold().replace(" ", "-")
        if not name and not sku:
            continue
        normalized_rows.append(
            {
                "sku": sku,
                "name": name or sku,
                "description": row.get(mapping.get("description", ""), "").strip(),
                "price": row.get(mapping.get("price", ""), "").strip(),
                "availability": row.get(mapping.get("availability", ""), "").strip(),
            }
        )
    return tuple(normalized_rows)


def parse_upload(
    filename: str,
    content: bytes,
    declared_mime: str | None,
    *,
    max_upload_bytes: int = 10_000_000,
    max_archive_bytes: int = 50_000_000,
    max_rows: int = 5_000,
) -> ParsedUpload:
    if not filename or len(content) > max_upload_bytes:
        raise Phase3Error("file_too_large", "The uploaded file is missing or too large.")
    suffix = _extension(filename)
    mime = (declared_mime or "").split(";", 1)[0].casefold()
    if suffix == ".svg" or mime == "image/svg+xml":
        safe = _svg(content)
        return ParsedUpload(safe, "vector", "image/svg+xml", {"sanitized": True})
    if suffix in {".png", ".jpg", ".jpeg", ".webp"} or mime.startswith("image/"):
        metadata, safe = _image(filename, content)
        expected = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}.get(suffix)
        if expected and metadata["format"] == "jpeg" and expected != "image/jpeg":
            raise Phase3Error("mime_mismatch", "The file extension does not match its content.")
        return ParsedUpload(safe, "image", Image.MIME.get(metadata["format"].upper(), mime), metadata)
    if suffix == ".pdf" or mime == "application/pdf":
        if not content.startswith(b"%PDF"):
            raise Phase3Error("mime_mismatch", "The file is not a valid PDF.")
        try:
            reader = PdfReader(BytesIO(content), strict=True)
            if len(reader.pages) > 100:
                raise Phase3Error("too_many_pages", "PDF contains too many pages.")
            extracted = "\n".join((page.extract_text() or "") for page in reader.pages)[:1_000_000]
        except Phase3Error:
            raise
        except Exception as exc:
            raise Phase3Error("malformed_pdf", "The PDF is malformed or encrypted.") from exc
        return ParsedUpload(content, "document", "application/pdf", {"pages": len(reader.pages)}, extracted)
    if suffix == ".docx" or mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        archive = _zip_guard(content, max_archive_bytes)
        try:
            names = {member.filename for member in archive.infolist()}
            if "word/document.xml" not in names or any("vbaProject.bin" in name for name in names):
                raise Phase3Error("unsafe_docx", "The DOCX contains unsupported active content.")
        finally:
            archive.close()
        try:
            document = Document(BytesIO(content))
            extracted = "\n".join(paragraph.text for paragraph in document.paragraphs)[:1_000_000]
        except Exception as exc:
            raise Phase3Error("malformed_docx", "The DOCX is malformed.") from exc
        return ParsedUpload(content, "document", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", {}, extracted)
    if suffix == ".xlsx" or mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        archive = _zip_guard(content, max_archive_bytes)
        try:
            names = {member.filename for member in archive.infolist()}
            if any("vbaProject.bin" in name or "externalLink" in name for name in names):
                raise Phase3Error("unsafe_xlsx", "The workbook contains unsupported active content.")
        finally:
            archive.close()
        try:
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            values = list(sheet.iter_rows(values_only=True, max_row=max_rows + 1))
            workbook.close()
        except Exception as exc:
            raise Phase3Error("malformed_xlsx", "The workbook is malformed.") from exc
        if not values:
            raise Phase3Error("missing_headers", "The workbook requires a header row.")
        headers = [str(value or "").strip() for value in values[0]]
        rows = tuple({headers[i]: str(value or "") for i, value in enumerate(row) if i < len(headers) and headers[i]} for row in values[1:])
        products = _product_rows(headers, rows)
        return ParsedUpload(content, "spreadsheet", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", {"columns": headers, "row_count": len(rows)}, product_rows=products)
    if suffix == ".csv" or mime in {"text/csv", "application/csv"}:
        headers, rows = _csv_rows(content, max_rows)
        products = _product_rows(headers, rows)
        return ParsedUpload(content, "csv", "text/csv", {"columns": headers, "row_count": len(rows)}, product_rows=products)
    raise Phase3Error("unsupported_file", "This file type is not supported.")


def _public_ip(address: str) -> bool:
    try:
        value = ipaddress.ip_address(address)
    except ValueError:
        return False
    return not (
        value.is_private
        or value.is_loopback
        or value.is_link_local
        or value.is_reserved
        or value.is_multicast
        or value.is_unspecified
    )


def validate_public_url(
    value: str,
    *,
    resolver: Callable[[str], list[str]] | None = None,
    resolve: bool = True,
) -> str:
    parsed = urlparse(value.strip())
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise Phase3Error("unsafe_url", "Only absolute HTTP(S) URLs are allowed.")
    if parsed.username or parsed.password:
        raise Phase3Error("unsafe_url", "URLs containing credentials are not allowed.")
    host = parsed.hostname.casefold()
    if host in {"localhost", "localhost.localdomain", "metadata.google.internal"} or host.endswith(".local"):
        raise Phase3Error("private_url", "Private and local destinations are not allowed.")
    try:
        literal = ipaddress.ip_address(host)
    except ValueError:
        literal = None
    if literal is not None and not _public_ip(str(literal)):
        raise Phase3Error("private_url", "Private and local destinations are not allowed.")
    if resolve:
        lookup = resolver or (lambda hostname: [item[4][0] for item in socket.getaddrinfo(hostname, None)])
        try:
            addresses = set(lookup(host))
        except OSError as exc:
            raise Phase3Error("dns_failed", "The URL hostname could not be resolved safely.") from exc
        if not addresses or any(not _public_ip(address) for address in addresses):
            raise Phase3Error("private_url", "The URL resolves to a private or local destination.")
    return parsed.geturl()


class Phase3Service:
    def __init__(
        self,
        sessions: sessionmaker[Session],
        accounts: AccountService,
        storage: LocalStorage,
        *,
        max_upload_bytes: int,
        max_archive_bytes: int,
        max_rows: int,
        offline: bool,
    ) -> None:
        self._sessions = sessions
        self._accounts = accounts
        self._storage = storage
        self._max_upload_bytes = max_upload_bytes
        self._max_archive_bytes = max_archive_bytes
        self._max_rows = max_rows
        self._offline = offline

    def upload_asset(
        self, principal: Principal, workspace_id: str, filename: str, content: bytes, mime: str | None,
        *, purpose: str | None = None
    ) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        parsed = parse_upload(
            filename,
            content,
            mime,
            max_upload_bytes=self._max_upload_bytes,
            max_archive_bytes=self._max_archive_bytes,
            max_rows=self._max_rows,
        )
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            digest = hashlib.sha256(parsed.content).hexdigest()
            from .asset_roles import asset_purpose, upload_roles
            purpose = purpose or ("product" if parsed.mime_type.startswith("image/") else "document")
            if purpose not in {"logo", "reference", "product", "document"}:
                raise Phase3Error("invalid_role", "Choose an upload purpose.")
            roles = upload_roles(session, workspace_id)
            existing = next((asset for asset in session.scalars(
                select(WorkspaceAsset).where(
                    WorkspaceAsset.workspace_id == workspace_id,
                    WorkspaceAsset.sha256 == digest,
                )
            ) if asset_purpose(asset, roles) == purpose), None)
            if existing is not None:
                existing.metadata_json = {**existing.metadata_json, "purpose": purpose}
                return self._asset_payload(existing)
            asset = WorkspaceAsset(
                workspace_id=workspace_id,
                storage_key="pending",
                original_name=Path(filename).name[:255],
                mime_type=parsed.mime_type,
                sha256=digest,
                size=len(parsed.content),
                asset_type=parsed.asset_type,
                status="ready",
                source="upload",
                width=parsed.metadata.get("width"),
                height=parsed.metadata.get("height"),
                orientation=parsed.metadata.get("orientation"),
                extracted_text=parsed.extracted_text,
                metadata_json={**parsed.metadata, "purpose": purpose},
            )
            session.add(asset)
            session.flush()
            asset.storage_key = f"workspaces/{workspace_id}/assets/{asset.id}/original{_extension(filename)}"
            self._storage.put_bytes(asset.storage_key, parsed.content)
            return self._asset_payload(asset)

    def create_import(
        self,
        principal: Principal,
        workspace_id: str,
        filename: str,
        content: bytes,
        mime: str | None,
        dedupe_key: str,
    ) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        parsed = parse_upload(
            filename,
            content,
            mime,
            max_upload_bytes=self._max_upload_bytes,
            max_archive_bytes=self._max_archive_bytes,
            max_rows=self._max_rows,
        )
        if not parsed.product_rows:
            raise Phase3Error("empty_import", "The import contains no usable product rows.")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            existing = session.scalar(select(ImportJob).where(ImportJob.workspace_id == workspace_id, ImportJob.dedupe_key == dedupe_key))
            if existing is not None:
                return self._import_payload(existing)
            job = ImportJob(
                workspace_id=workspace_id,
                created_by_user_id=principal.user_id,
                kind="product_file",
                status="preview_ready",
                source_name=Path(filename).name[:255],
                dedupe_key=dedupe_key,
                preview={"columns": parsed.metadata.get("columns", []), "rows": list(parsed.product_rows)},
                checkpoint={},
            )
            session.add(job)
            session.flush()
            key = f"workspaces/{workspace_id}/imports/{job.id}/source{_extension(filename)}"
            self._storage.put_bytes(key, content)
            job.checkpoint = {"storage_key": key, "row_count": len(parsed.product_rows)}
            return self._import_payload(job)

    def create_url_import(self, principal: Principal, workspace_id: str, url: str, dedupe_key: str) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        normalized = validate_public_url(url, resolve=not self._offline)
        status = "blocked_offline" if self._offline else "queued"
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            existing = session.scalar(select(ImportJob).where(ImportJob.workspace_id == workspace_id, ImportJob.dedupe_key == dedupe_key))
            if existing is not None:
                return self._import_payload(existing)
            job = ImportJob(
                workspace_id=workspace_id,
                created_by_user_id=principal.user_id,
                kind="website",
                status=status,
                source_url=normalized,
                dedupe_key=dedupe_key,
                preview={},
                checkpoint={},
                error_code="external_disabled" if self._offline else None,
                error_message="Website imports are disabled in offline_test mode." if self._offline else None,
            )
            session.add(job)
            session.flush()
            return self._import_payload(job)

    def get_import(self, principal: Principal, workspace_id: str, import_id: str) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "read")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            job = session.scalar(select(ImportJob).where(ImportJob.id == import_id, ImportJob.workspace_id == workspace_id))
            if job is None:
                raise Phase3Error("import_not_found", "Import was not found.")
            return self._import_payload(job)

    def list_assets(self, principal: Principal, workspace_id: str) -> list[dict[str, Any]]:
        self._accounts.require_permission(principal, workspace_id, "read")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            rows = session.scalars(
                select(WorkspaceAsset)
                .where(WorkspaceAsset.workspace_id == workspace_id)
                .order_by(WorkspaceAsset.created_at.desc(), WorkspaceAsset.id)
            )
            from .asset_roles import asset_purpose, upload_roles
            roles = upload_roles(session, workspace_id)
            return [{**self._asset_payload(asset), "purpose": asset_purpose(asset, roles)} for asset in rows]

    def confirm_import(self, principal: Principal, workspace_id: str, import_id: str) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            job = session.scalar(select(ImportJob).where(ImportJob.id == import_id, ImportJob.workspace_id == workspace_id))
            if job is None:
                raise Phase3Error("import_not_found", "Import was not found.")
            if job.kind != "product_file":
                raise Phase3Error("import_not_found", "Import was not found.")
            if job.status == "confirmed":
                return {"import_id": job.id, "status": job.status, "products_created": job.checkpoint.get("products_created", 0)}
            created = 0
            for row in job.preview.get("rows", []):
                product = session.scalar(select(Product).where(Product.workspace_id == workspace_id, Product.sku == row["sku"]))
                if product is None:
                    product = Product(workspace_id=workspace_id, sku=row["sku"], name=row["name"])
                    session.add(product)
                    created += 1
                product.name = row["name"]
                product.description = row.get("description") or None
                product.price = row.get("price") or None
                product.availability = row.get("availability") or None
                product.source_ref = job.id
            job.status = "confirmed"
            job.completed_at = utc_now()
            job.checkpoint = {**job.checkpoint, "products_created": created}
            return {"import_id": job.id, "status": job.status, "products_created": created}

    def profile_proposal(self, principal: Principal, workspace_id: str, fields: dict[str, Any]) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        if not fields:
            raise Phase3Error("empty_profile", "At least one profile field is required.")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            latest = session.scalar(select(func.max(BrandProfileVersion.version)).where(BrandProfileVersion.workspace_id == workspace_id)) or 0
            profile = BrandProfileVersion(workspace_id=workspace_id, version=latest + 1, status="draft", fields=fields, provenance={"source": "owner_manual"}, created_by_user_id=principal.user_id)
            session.add(profile)
            session.flush()
            return self._profile_payload(profile)

    def confirm_profile(self, principal: Principal, workspace_id: str, version: int) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "approve")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            profile = session.scalar(select(BrandProfileVersion).where(BrandProfileVersion.workspace_id == workspace_id, BrandProfileVersion.version == version))
            if profile is None:
                raise Phase3Error("profile_not_found", "Brand profile version was not found.")
            profile.status = "confirmed"
            profile.confirmed_at = utc_now()
            return self._profile_payload(profile)

    def current_profile(self, principal: Principal, workspace_id: str) -> dict[str, Any] | None:
        self._accounts.require_permission(principal, workspace_id, "read")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            profile = session.scalar(select(BrandProfileVersion).where(BrandProfileVersion.workspace_id == workspace_id, BrandProfileVersion.status == "confirmed").order_by(BrandProfileVersion.version.desc()))
            return self._profile_payload(profile) if profile else None

    def onboarding_status(self, principal: Principal, workspace_id: str) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "read")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            interview = self._interview(session, workspace_id)
            next_key = self._next_question(interview)
            return self._interview_payload(interview, next_key)

    def onboarding_turn(self, principal: Principal, workspace_id: str, question_key: str, answer: str, confirmed: bool) -> dict[str, Any]:
        self._accounts.require_permission(principal, workspace_id, "write")
        if not answer.strip():
            raise Phase3Error("empty_answer", "An answer is required.")
        with self._sessions() as session, session.begin():
            set_workspace_context(session, workspace_id)
            interview = self._interview(session, workspace_id)
            existing = interview.confirmed_answers.get(question_key)
            contradiction = None
            if existing is not None and existing != answer.strip():
                contradiction = {"question_key": question_key, "previous": existing, "new": answer.strip()}
                interview.contradictions = [*(interview.contradictions or []), contradiction]
            interview.asked_keys = list(dict.fromkeys([*(interview.asked_keys or []), question_key]))
            if confirmed:
                interview.confirmed_answers = {**(interview.confirmed_answers or {}), question_key: answer.strip()}
            turn = session.scalar(select(OnboardingTurn).where(OnboardingTurn.interview_id == interview.id, OnboardingTurn.question_key == question_key))
            if turn is None:
                session.add(OnboardingTurn(interview_id=interview.id, workspace_id=workspace_id, question_key=question_key, answer=answer.strip(), confirmed=confirmed))
            else:
                turn.answer = answer.strip()
                turn.confirmed = confirmed
            next_key = self._next_question(interview)
            if next_key is None:
                interview.status = "complete"
            return {"status": interview.status, "next_question": next_key, "contradiction": contradiction}

    @staticmethod
    def _interview(session: Session, workspace_id: str) -> OnboardingInterview:
        interview = session.scalar(select(OnboardingInterview).where(OnboardingInterview.workspace_id == workspace_id).order_by(OnboardingInterview.created_at))
        if interview is None:
            interview = OnboardingInterview(workspace_id=workspace_id, confirmed_answers={}, asked_keys=[], contradictions=[])
            session.add(interview)
            session.flush()
        return interview

    @staticmethod
    def _next_question(interview: OnboardingInterview) -> str | None:
        for key in ("brand_name", "category", "audience", "location", "offer", "price", "address"):
            if key not in (interview.confirmed_answers or {}):
                return key
        return None

    @staticmethod
    def _asset_payload(asset: WorkspaceAsset) -> dict[str, Any]:
        from .asset_roles import asset_purpose
        return {"id": asset.id, "workspace_id": asset.workspace_id, "name": asset.original_name, "mime_type": asset.mime_type, "sha256": asset.sha256, "size": asset.size, "asset_type": asset.asset_type, "source": asset.source, "status": asset.status, "metadata": asset.metadata_json, "purpose": asset_purpose(asset)}

    @staticmethod
    def _import_payload(job: ImportJob) -> dict[str, Any]:
        return {"id": job.id, "kind": job.kind, "status": job.status, "source_name": job.source_name, "source_url": job.source_url, "preview": job.preview, "error_code": job.error_code, "error_message": job.error_message}

    @staticmethod
    def _profile_payload(profile: BrandProfileVersion | None) -> dict[str, Any] | None:
        if profile is None:
            return None
        return {"id": profile.id, "version": profile.version, "status": profile.status, "fields": profile.fields, "provenance": profile.provenance, "confirmed_at": profile.confirmed_at.isoformat() if profile.confirmed_at else None}

    @staticmethod
    def _interview_payload(interview: OnboardingInterview, next_key: str | None) -> dict[str, Any]:
        return {"id": interview.id, "status": interview.status, "confirmed_answers": interview.confirmed_answers, "asked_keys": interview.asked_keys, "contradictions": interview.contradictions, "next_question": next_key}
